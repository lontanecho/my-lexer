from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "实验报告.md"
OUTPUT = ROOT / "词法分析器实验报告.docx"


def set_run_font(run, name="宋体", size=12, bold=None, color="000000",
                 east_asia=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    for slot in ("ascii", "hAnsi", "cs"):
        r_fonts.set(qn(f"w:{slot}"), name)
    r_fonts.set(qn("w:eastAsia"), east_asia or name)


def set_paragraph(paragraph, size=12, line=1.5, first_line=True,
                  before=0, after=0, keep=False):
    fmt = paragraph.paragraph_format
    fmt.line_spacing = line
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.keep_with_next = keep
    if first_line:
        fmt.first_line_indent = Cm(0.85)
    for run in paragraph.runs:
        set_run_font(run, size=size)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start),
                          ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "6")
        tag.set(qn("w:color"), "D9D9D9")


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


INLINE_PATTERN = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*)")


def add_inline(paragraph, text, size=12, color="000000"):
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=size - 0.5, color=color,
                         east_asia="宋体")
        else:
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True, color=color)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color)


def add_heading(document, text, level):
    if level == 0:
        paragraph = document.add_paragraph(style="Title")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(36)
        paragraph.paragraph_format.space_after = Pt(24)
        run = paragraph.add_run(text)
        set_run_font(run, size=18, bold=True)
        return
    style = "Heading 1" if level == 1 else "Heading 2"
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    paragraph.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=14 if level == 1 else 12, bold=True)


def add_body(document, text):
    paragraph = document.add_paragraph()
    add_inline(paragraph, text)
    set_paragraph(paragraph)


def add_list_item(document, text, ordered):
    paragraph = document.add_paragraph(style="List Number" if ordered else "List Bullet")
    paragraph.paragraph_format.left_indent = Cm(0.74)
    paragraph.paragraph_format.first_line_indent = Cm(-0.42)
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(0)
    add_inline(paragraph, text)


def add_code_block(document, code):
    for line in code.splitlines() or [""]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.6)
        paragraph.paragraph_format.right_indent = Cm(0.6)
        paragraph.paragraph_format.line_spacing = 1.05
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        p_pr = paragraph._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F3F5F7")
        p_pr.append(shd)
        run = paragraph.add_run(line if line else " ")
        set_run_font(run, name="Consolas", size=9.5, east_asia="宋体")


def add_table(document, rows):
    if not rows:
        return
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for row_index, values in enumerate(rows):
        row = table.rows[0] if row_index == 0 else table.add_row()
        if row_index == 0:
            repeat_table_header(row)
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            shade_cell(cell, "1F4E78" if row_index == 0 else
                       ("EEF3F8" if row_index % 2 == 0 else "FFFFFF"))
            paragraph = cell.paragraphs[0]
            paragraph.alignment = (WD_ALIGN_PARAGRAPH.CENTER if
                                   column_index < 3 and len(values) > 2 else
                                   WD_ALIGN_PARAGRAPH.LEFT)
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_before = Pt(1)
            paragraph.paragraph_format.space_after = Pt(1)
            add_inline(paragraph, value, size=12,
                       color="FFFFFF" if row_index == 0 else "000000")
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    spacer.paragraph_format.line_spacing = 0.5


def parse_table(lines, start):
    def cells(line):
        return [cell.strip() for cell in line.strip().strip("|").split("|")]
    rows = [cells(lines[start])]
    index = start + 2
    while index < len(lines) and lines[index].strip().startswith("|"):
        rows.append(cells(lines[index]))
        index += 1
    return rows, index


def build_document(markdown):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.5)

    normal = document.styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_run = header.add_run("编译原理实验报告")
    set_run_font(header_run, size=9)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page_run = footer.add_run()
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    page_run._r.extend([fld_char_1, instr_text, fld_char_2])
    set_run_font(page_run, size=9)

    lines = markdown.splitlines()
    index = 0
    paragraph_buffer = []

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            add_body(document, " ".join(part.strip() for part in paragraph_buffer))
            paragraph_buffer = []

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        if stripped.startswith("```"):
            flush_paragraph()
            index += 1
            code_lines = []
            while index < len(lines) and not lines[index].strip().startswith("```"):
                code_lines.append(lines[index])
                index += 1
            add_code_block(document, "\n".join(code_lines))
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and re.match(
                r"^\s*\|?[\s:|-]+\|", lines[index + 1]):
            flush_paragraph()
            rows, index = parse_table(lines, index)
            add_table(document, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1)) - 1
            add_heading(document, heading.group(2), level)
            index += 1
            continue
        ordered = re.match(r"^\d+\.\s+(.+)$", stripped)
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if ordered or bullet:
            flush_paragraph()
            add_list_item(document, (ordered or bullet).group(1), bool(ordered))
            index += 1
            continue
        paragraph_buffer.append(stripped)
        index += 1
    flush_paragraph()
    return document


if __name__ == "__main__":
    report = build_document(SOURCE.read_text(encoding="utf-8"))
    report.core_properties.title = "C/C++ 源程序词法分析器设计与实现实验报告"
    report.core_properties.subject = "编译原理词法分析实验"
    report.core_properties.author = ""
    report.save(OUTPUT)
    print(OUTPUT)
